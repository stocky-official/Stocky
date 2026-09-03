"""
Fast Deep analysis of data/Supplier.xlsx using read_only mode
Rule 2 Mandate: Resides strictly inside _technical_support/scripts/
"""

import openpyxl
import os

FILEPATH = r"c:\Users\abdelrahman.mamdouh_\Desktop\Stocky\data\Supplier.xlsx"

def inspect_supplier_file():
    print(f"Loading {FILEPATH} in read_only mode...")
    wb = openpyxl.load_workbook(FILEPATH, read_only=True, data_only=True)
    print(f"Sheets found: {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n==================================================")
        print(f"SHEET: '{sheet_name}'")
        print(f"==================================================")

        row_count = 0
        non_empty_rows = 0
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            row_count += 1
            # Check if non-empty
            if any(cell is not None for cell in row):
                non_empty_rows += 1
                if non_empty_rows <= 12:
                    clean_row = [str(cell)[:50] if cell is not None else None for cell in row[:10]]
                    while clean_row and clean_row[-1] is None:
                        clean_row.pop()
                    print(f"Row {i:4d}: {clean_row}")

        print(f"-> Total scanned rows: {row_count:,}, Non-empty rows: {non_empty_rows:,}")

    wb.close()

if __name__ == "__main__":
    inspect_supplier_file()
