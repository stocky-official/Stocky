import openpyxl
import os
import glob

data_dir = r"c:\Users\abdelrahman.mamdouh_\Desktop\Stocky\data"
files = [
    "Cookies.xlsx",
    "Crispy Chicken.xlsx",
    "Edible Grocery.xlsx",
    "AutoMotive.xlsx"
]

for filename in files:
    filepath = os.path.join(data_dir, filename)
    if not os.path.exists(filepath):
        continue
    print(f"\n==================== {filename} ====================")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    print("Sheets:", wb.sheetnames)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    print(f"Total rows: {len(rows)}")
    for i, row in enumerate(rows[:6]):
        print(f"Row {i}: {row}")
