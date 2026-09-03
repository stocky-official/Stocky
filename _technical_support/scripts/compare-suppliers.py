"""
Analyze name overlap between 'Suppler' directory and 'Sheet1' items
Rule 2 Mandate: Resides strictly in _technical_support/scripts/
"""

import openpyxl

FILEPATH = r"c:\Users\abdelrahman.mamdouh_\Desktop\Stocky\data\Supplier.xlsx"

wb = openpyxl.load_workbook(FILEPATH, read_only=True, data_only=True)

# 1. Directory names
dir_suppliers = set()
for row in wb["Suppler"].iter_rows(values_only=True):
    if row[0] and str(row[0]).strip().lower() != "company  name":
        dir_suppliers.add(str(row[0]).strip())

# 2. Sheet1 names
item_suppliers = set()
for i, row in enumerate(wb["Sheet1"].iter_rows(values_only=True)):
    if i > 0 and row[2]:
        item_suppliers.add(str(row[2]).strip())

wb.close()

print(f"Directory suppliers count: {len(dir_suppliers)}")
print(f"Sheet1 suppliers count: {len(item_suppliers)}")

exact_matches = dir_suppliers.intersection(item_suppliers)
print(f"Exact name matches: {len(exact_matches)}")

in_dir_not_sheet1 = dir_suppliers - item_suppliers
in_sheet1_not_dir = item_suppliers - dir_suppliers

print(f"\nIn Directory but not in Sheet1 ({len(in_dir_not_sheet1)}):")
for s in sorted(in_dir_not_sheet1)[:10]:
    print(f" - {s}")

print(f"\nIn Sheet1 but not in Directory ({len(in_sheet1_not_dir)}):")
for s in sorted(in_sheet1_not_dir)[:10]:
    print(f" - {s}")
