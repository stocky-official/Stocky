"""
Deep Cleaning Analysis for data/Supplier.xlsx
Rule 2 Mandate: Resides strictly in _technical_support/scripts/
"""

import openpyxl
import re

FILEPATH = r"c:\Users\abdelrahman.mamdouh_\Desktop\Stocky\data\Supplier.xlsx"

def analyze_sheet_suppler():
    print("\n--- ANALYZING SHEET 'Suppler' (Directory) ---")
    wb = openpyxl.load_workbook(FILEPATH, read_only=True, data_only=True)
    sheet = wb["Suppler"]

    rows = list(sheet.iter_rows(values_only=True))
    wb.close()

    header = rows[0]
    print(f"Header: {header}")
    data_rows = rows[1:]
    print(f"Total data rows in 'Suppler': {len(data_rows)}")

    companies = {}
    formatting_issues = []

    for idx, row in enumerate(data_rows, start=2):
        raw_company = row[0]
        raw_category = row[1]
        raw_email = row[2]
        raw_person = row[3]
        raw_phone = row[4]

        if not raw_company or str(raw_company).strip() == "":
            formatting_issues.append((idx, "Empty company name", row))
            continue

        comp_name = str(raw_company).strip()

        # Check issues
        issues = []
        if raw_person is None and raw_phone is not None:
            issues.append("Person is None but phone exists")
        if raw_person and ":" in str(raw_person):
            issues.append("Colon in contact person (likely name:phone)")
        if raw_phone and ":" in str(raw_phone):
            issues.append("Colon in contact phone (likely name:phone)")
        if raw_email and (";" in str(raw_email) or "," in str(raw_email)):
            issues.append("Multiple emails in email field")
        if raw_phone and ("/" in str(raw_phone) or " " in str(raw_phone).strip()):
            issues.append("Phone number with spaces or multiple numbers")

        if comp_name not in companies:
            companies[comp_name] = []
        companies[comp_name].append({
            "row": idx,
            "category": raw_category,
            "email": raw_email,
            "person": raw_person,
            "phone": raw_phone,
            "issues": issues
        })

    print(f"\nDistinct company names: {len(companies)}")
    duplicate_companies = {k: v for k, v in companies.items() if len(v) > 1}
    print(f"Companies with multiple rows: {len(duplicate_companies)}")
    print("\nSample duplicate companies:")
    for comp, rows in list(duplicate_companies.items())[:5]:
        print(f" - '{comp}' ({len(rows)} rows)")
        for r in rows:
            print(f"     Row {r['row']}: Cat={r['category']} | Email={r['email']} | Person={r['person']} | Phone={r['phone']}")

    print("\nSample formatting issues detected:")
    sample_issues = 0
    for comp, rows in companies.items():
        for r in rows:
            if r["issues"] and sample_issues < 10:
                print(f" - [{comp}] Row {r['row']}: {', '.join(r['issues'])}")
                print(f"     Person: {r['person']} | Phone: {r['phone']} | Email: {r['email']}")
                sample_issues += 1

def analyze_sheet1_items():
    print("\n--- ANALYZING SHEET 'Sheet1' (Items per Supplier) ---")
    wb = openpyxl.load_workbook(FILEPATH, read_only=True, data_only=True)
    sheet = wb["Sheet1"]

    supplier_item_counts = {}
    rows_count = 0
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            continue
        rows_count += 1
        supplier = str(row[2]).strip() if row[2] else "NO_SUPPLIER"
        supplier_item_counts[supplier] = supplier_item_counts.get(supplier, 0) + 1

    wb.close()
    print(f"Total rows in Sheet1: {rows_count:,}")
    print(f"Distinct suppliers in Sheet1: {len(supplier_item_counts)}")
    print("Top 10 suppliers by item count in Sheet1:")
    for s, c in sorted(supplier_item_counts.items(), key=lambda x: x[1], reverse=True)[:10]:
        print(f" - {s}: {c:,} items")

if __name__ == "__main__":
    analyze_sheet_suppler()
    analyze_sheet1_items()
