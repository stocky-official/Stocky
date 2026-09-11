"""
Supplier Cleaning & Ingestion Script
Imports data/Supplier.xlsx into Supabase PostgreSQL (public.suppliers and public.supplier_items)
Rule 2 Mandate: Resides strictly in _technical_support/scripts/
"""

import os
import re
import uuid
import openpyxl
import psycopg2
from psycopg2.extras import execute_values

DB_URL = os.environ["DATABASE_URL"]
FILEPATH = r"c:\Users\abdelrahman.mamdouh_\Desktop\Stocky\data\Supplier.xlsx"

def clean_phone_str(raw_phone):
    if not raw_phone:
        return "—"
    val = str(raw_phone).strip()
    # Replace letter O / o with digit 0
    val = re.sub(r'\b[Oo](\d{8,11})\b', r'0\1', val)
    val = re.sub(r'^[Oo]', '0', val)
    # Remove excessive whitespace
    val = re.sub(r'\s+', ' ', val)
    return val if val else "—"

def clean_email_str(raw_email):
    if not raw_email:
        return None
    val = str(raw_email).strip().strip(';, ')
    # Split by semicolon, comma, or newline
    tokens = [e.strip().replace(' ', '') for e in re.split(r'[;,\n]', val) if e.strip()]
    valid_emails = [e for e in tokens if '@' in e and '.' in e]
    if valid_emails:
        return valid_emails[0] # primary valid email
    return tokens[0] if tokens else None

def parse_contact_and_phone(raw_person, raw_phone):
    person = str(raw_person).strip() if raw_person is not None else ""
    phone = str(raw_phone).strip() if raw_phone is not None else ""

    # Replace capital O with 0
    person = re.sub(r'\b[Oo](\d{8,11})\b', r'0\1', person)
    phone = re.sub(r'\b[Oo](\d{8,11})\b', r'0\1', phone)

    # Clean multiline: if multiline, take first line
    if '\n' in person:
        person = person.split('\n')[0].strip()
    if '\n' in phone:
        phone = phone.split('\n')[0].strip()

    # Case 1: person is blank/None, but phone has "Name: Phone"
    if (not person or person.lower() in ['none', 'nan', '']) and ":" in phone:
        parts = phone.split(":", 1)
        return parts[0].strip(), clean_phone_str(parts[1].strip())

    # Case 2: person has "Name: Phone"
    if ":" in person:
        parts = person.split(":", 1)
        cand_name = parts[0].strip()
        cand_phone = clean_phone_str(parts[1].strip())
        if phone and ":" in phone:
            phone = clean_phone_str(phone.split(":", 1)[1].strip())
        final_phone = phone if phone and phone != "—" else cand_phone
        return cand_name, final_phone

    # Case 3: phone has "Name: Phone"
    if ":" in phone:
        parts = phone.split(":", 1)
        name = person if person else parts[0].strip()
        final_phone = clean_phone_str(parts[1].strip())
        return name, final_phone

    final_name = person if person else "Procurement Contact"
    final_phone = clean_phone_str(phone)
    return final_name, final_phone

def run_supplier_seeding():
    print("==================================================")
    print("       STOCKY SUPPLIER DATA INGESTION             ")
    print("==================================================")

    conn = psycopg2.connect(DB_URL, sslmode="require")
    cur = conn.cursor()

    try:
        # 1. Get Circle K Company ID
        cur.execute("SELECT id, name FROM public.companies WHERE code = 'CRK' LIMIT 1;")
        comp_row = cur.fetchone()
        if not comp_row:
            raise Exception("Circle K company not found in database!")
        company_id = comp_row[0]
        print(f"Target company: {comp_row[1]} (ID: {company_id})")

        # 2. Clear old supplier records
        print("\n[1/5] Purging previous supplier tables...")
        cur.execute("DELETE FROM public.supplier_items;")
        cur.execute("DELETE FROM public.suppliers;")
        conn.commit()
        print(" -- Tables cleared.")

        # 3. Load Sheet2 Normalization Mappings
        print("\n[2/5] Reading Sheet2 aliases...")
        wb = openpyxl.load_workbook(FILEPATH, read_only=True, data_only=True)
        mappings = {}
        for row in wb["Sheet2"].iter_rows(values_only=True):
            if row[0] and row[1]:
                alias = str(row[0]).strip().lower()
                canonical = str(row[1]).strip()
                if alias != "company  name":
                    mappings[alias] = canonical
        print(f" -- Loaded {len(mappings)} alias mappings from Sheet2.")

        # 4. Parse 'Suppler' Directory Sheet
        print("\n[3/5] Parsing and cleaning 'Suppler' directory...")
        suppliers_dict = {} # canonical_name -> dict

        for i, row in enumerate(wb["Suppler"].iter_rows(values_only=True)):
            if i == 0 or not row[0]:
                continue
            raw_comp = str(row[0]).strip()
            if not raw_comp:
                continue

            norm_key = raw_comp.lower()
            canonical_name = mappings.get(norm_key, re.sub(r'\s+', ' ', raw_comp))

            c_name, c_phone = parse_contact_and_phone(row[3], row[4])
            c_email = clean_email_str(row[2])

            raw_cat = str(row[1]).strip() if row[1] else ""
            cat_set = set()
            if raw_cat:
                for c in raw_cat.split(','):
                    if c.strip():
                        cat_set.add(c.strip())

            if canonical_name not in suppliers_dict:
                suppliers_dict[canonical_name] = {
                    "id": str(uuid.uuid4()),
                    "name": canonical_name,
                    "contact_name": c_name,
                    "contact_phone": c_phone,
                    "contact_email": c_email,
                    "categories": cat_set,
                    "items": []
                }
            else:
                # Merge categories
                suppliers_dict[canonical_name]["categories"].update(cat_set)
                # If existing has default or empty contact, override with more complete one
                curr = suppliers_dict[canonical_name]
                if (curr["contact_phone"] == "—" or curr["contact_name"] == "Procurement Contact") and c_phone != "—":
                    curr["contact_name"] = c_name
                    curr["contact_phone"] = c_phone
                if not curr["contact_email"] and c_email:
                    curr["contact_email"] = c_email

        print(f" -- Processed {len(suppliers_dict)} distinct directory suppliers.")

        # 5. Parse 'Sheet1' Catalog (15,919 items)
        print("\n[4/5] Parsing and linking 15,919 catalog items from Sheet1...")
        item_records = [] # (id, supplier_id, item_name, item_category, notes)

        for i, row in enumerate(wb["Sheet1"].iter_rows(values_only=True)):
            if i == 0 or not row[2]:
                continue

            raw_sup = str(row[2]).strip()
            norm_key = raw_sup.lower()
            canonical_name = mappings.get(norm_key, re.sub(r'\s+', ' ', raw_sup))

            barcode = str(row[0]).strip() if row[0] is not None else None
            if barcode and barcode.endswith('.0'):
                barcode = barcode[:-2]
            dept = str(row[1]).strip() if row[1] else None
            desc = str(row[3]).strip() if row[3] else "Item"

            # Check if supplier exists, else create from Sheet1 contact
            if canonical_name not in suppliers_dict:
                c_name, c_phone = parse_contact_and_phone(row[5], row[6])
                c_email = clean_email_str(row[4])
                suppliers_dict[canonical_name] = {
                    "id": str(uuid.uuid4()),
                    "name": canonical_name,
                    "contact_name": c_name,
                    "contact_phone": c_phone,
                    "contact_email": c_email,
                    "categories": set([dept]) if dept else set(),
                    "items": []
                }

            sup_id = suppliers_dict[canonical_name]["id"]
            if dept:
                suppliers_dict[canonical_name]["categories"].add(dept)

            item_id = str(uuid.uuid4())
            item_records.append((
                item_id,
                sup_id,
                desc[:250],
                dept[:150] if dept else None,
                f"Barcode: {barcode}" if barcode else None
            ))

        wb.close()
        print(f" -- Linked {len(item_records):,} products across {len(suppliers_dict)} suppliers.")

        # 6. Insert Suppliers into public.suppliers
        print(f"\n[5/5] Ingesting {len(suppliers_dict)} suppliers into database...")
        sup_rows = [
            (
                s["id"],
                company_id,
                s["name"][:250],
                s["contact_name"][:250],
                s["contact_phone"][:50],
                s["contact_email"][:250] if s["contact_email"] else None
            )
            for s in suppliers_dict.values()
        ]

        execute_values(
            cur,
            """
            INSERT INTO public.suppliers (id, company_id, name, contact_name, contact_phone, contact_email)
            VALUES %s;
            """,
            sup_rows
        )
        conn.commit()
        print(f" -- {len(sup_rows)} suppliers inserted.")

        # 7. Insert Supplier Items in batches
        print(f" -- Ingesting {len(item_records):,} supplier items in batches...")
        chunk_size = 5000
        for i in range(0, len(item_records), chunk_size):
            chunk = item_records[i : i + chunk_size]
            execute_values(
                cur,
                """
                INSERT INTO public.supplier_items (id, supplier_id, item_name, item_category, notes)
                VALUES %s;
                """,
                chunk
            )
            conn.commit()
            print(f"    Inserted {min(i + chunk_size, len(item_records)):,} / {len(item_records):,} items...")

        print("\n==================================================")
        print("SUCCESS: Supplier directory & items seeded cleanly!")
        print("==================================================")

    except Exception as e:
        conn.rollback()
        print(f"Error during supplier seeding: {e}")
        raise e
    finally:
        cur.close()
        conn.close()

if __name__ == "__main__":
    run_supplier_seeding()
